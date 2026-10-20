# md_to_docx.py — convert the PRD markdown into a Word document
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

src, dst = sys.argv[1], sys.argv[2]
lines = open(src, encoding='utf-8').read().splitlines()
doc = Document()

style = doc.styles['Normal']
style.font.name = 'Microsoft YaHei'
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

for i in range(1, 4):
    h = doc.styles['Heading %d' % i]
    h.font.name = 'Microsoft YaHei'
    h.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

def add_runs(par, text):
    parts = re.split(r'(\*\*.+?\*\*|`[^`]+`)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith('`') and part.endswith('`'):
            r = par.add_run(part[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(9.5)
            r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
        else:
            par.add_run(part)

i = 0
while i < len(lines):
    line = lines[i]
    if not line.strip():
        i += 1; continue
    if line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-{2,}:?', c) for c in cells):
                rows.append(cells)
            i += 1
        if rows:
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = 'Table Grid'
            for r, row in enumerate(rows):
                for c, cell in enumerate(row[:len(rows[0])]):
                    p = table.cell(r, c).paragraphs[0]
                    add_runs(p, cell)
                    if r == 0:
                        for run in p.runs:
                            run.bold = True
        continue
    m = re.match(r'^(#{1,3})\s+(.*)$', line)
    if m:
        level = len(m.group(1)) - 1
        doc.add_heading(m.group(2), level=level)
        i += 1; continue
    if re.match(r'^[-*]\s+', line):
        p = doc.add_paragraph(style='List Bullet')
        add_runs(p, re.sub(r'^[-*]\s+', '', line))
        i += 1; continue
    if re.match(r'^\d+\.\s+', line):
        p = doc.add_paragraph(style='List Number')
        add_runs(p, re.sub(r'^\d+\.\s+', '', line))
        i += 1; continue
    p = doc.add_paragraph()
    add_runs(p, line)
    i += 1

doc.save(dst)
print('saved', dst)
